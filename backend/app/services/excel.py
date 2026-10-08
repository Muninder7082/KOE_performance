"""ExcelReportService — maintains ONE ever-growing workbook.

Guarantees
  * "Performance History" is append-only. Existing rows are never edited,
    re-ordered or removed. Every PerformanceResult is written exactly once
    (tracked with performance_results.excel_appended + a Result ID column,
    so a crash between "upload" and "mark appended" cannot duplicate rows).
  * If the workbook is missing it is rebuilt from PostgreSQL, which holds the
    complete history, so the new file again contains ALL records.
  * A workbook that cannot be opened is never overwritten: it is preserved
    as website-performance.corrupt-<timestamp>.xlsx and a complete
    replacement is rebuilt from the database.
  * After saving, the file is re-opened and the history row count must be
    >= the previous count before it replaces the stored copy.
  * Updates are serialised by an asyncio lock (same process) and a database
    lease lock (other processes/containers).
The three other sheets are snapshots and are regenerated on every update.
"""
from __future__ import annotations

import asyncio
import logging
import tempfile
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from .. import locks
from ..clock import now_utc
from ..config import get_settings
from ..models import MonitoredWebsite, PerformanceResult, ScheduledTask
from . import global_schedule, reporting, settings_service
from .pagespeed import psi_report_link
from .storage import StorageService

log = logging.getLogger(__name__)

HISTORY = "Performance History"
LATEST = "Latest Status"
ATTENTION = "Attention Required"
SUMMARY = "Monitoring Summary"
HISTORY_HEADERS = [
    "Date", "Time", "Website Name", "URL", "Device", "Performance Score", "Accessibility Score",
    "Best Practices Score", "SEO Score", "FCP", "LCP", "TBT", "CLS", "Speed Index", "Status",
    "Final URL", "Test Status", "Error", "Result ID",
]
RESULT_ID_COL = len(HISTORY_HEADERS)  # 1-based index of "Result ID"
MAX_SHEET_ROWS = 1_048_576  # Excel's hard limit; history rolls over to "Performance History 2", ...
LOCK_NAME = "excel_workbook"

_HEADER_FILL = PatternFill("solid", fgColor="0F2B46")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_FILLS = {
    "GOOD": PatternFill("solid", fgColor="D1FAE5"),
    "ATTENTION": PatternFill("solid", fgColor="FEF3C7"),
    "FAILED": PatternFill("solid", fgColor="FEE2E2"),
}
_FORMATS = {10: '0.00" s"', 11: '0.00" s"', 12: '0" ms"', 13: "0.000", 14: '0.00" s"'}
_WIDTHS = [13, 8, 26, 48, 9, 12, 13, 14, 9, 9, 9, 10, 8, 12, 11, 48, 11, 50, 10]

_process_lock = asyncio.Lock()


class ExcelError(Exception):
    pass


@dataclass
class SyncResult:
    appended: int
    rebuilt: bool
    total_rows: int
    preserved_corrupt_copy: str | None = None


def _status(r: PerformanceResult) -> str:
    return reporting.result_status(r)


def history_row(r: PerformanceResult, tz: ZoneInfo) -> list:
    local = r.tested_at.astimezone(tz)
    error = None
    if r.status != "success":
        error = f"{r.error_code or 'error'}: {r.error_message or ''}".strip()[:1000]
    return [
        local.date(), local.time().replace(second=0, microsecond=0), r.website_name, r.requested_url,
        r.strategy.capitalize(), r.performance_score, r.accessibility_score, r.best_practices_score,
        r.seo_score, r.fcp_s, r.lcp_s, r.tbt_ms, r.cls, r.speed_index_s, _status(r), r.final_url,
        "SUCCESS" if r.status == "success" else "FAILED", error, r.id,
    ]


def _clean(v):
    """Strip control characters openpyxl/Excel cannot store (they would make every save fail)."""
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def _put(ws, row: list) -> None:
    ws.append([_clean(v) for v in row])


def _style_header(ws, headers: list[str], widths: list[int] | None = None) -> None:
    _put(ws, headers)
    for i, _ in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = _HEADER_FILL, _HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = (widths or [])[i - 1] if widths and i <= len(widths) else 16
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def _append_history(ws, row: list) -> None:
    _put(ws, row)
    r = ws.max_row
    ws.cell(r, 1).number_format = "DD-MMM-YYYY"
    ws.cell(r, 2).number_format = "HH:MM"
    for col, fmt in _FORMATS.items():
        ws.cell(r, col).number_format = fmt
    fill = _FILLS.get(row[14])
    if fill:
        ws.cell(r, 6).fill = fill
        ws.cell(r, 15).fill = fill


@dataclass
class Snapshot:
    tz_name: str
    generated_at: datetime
    websites: list[dict]
    attention: list[dict]
    daily: list[dict]
    totals: dict


def _write_snapshot_sheets(wb: Workbook, snap: Snapshot) -> None:
    for name in (LATEST, ATTENTION, SUMMARY):
        if name in wb.sheetnames:
            del wb[name]
    tz = ZoneInfo(snap.tz_name)
    gen = snap.generated_at.astimezone(tz).strftime("%d-%b-%Y %H:%M")

    ws = wb.create_sheet(LATEST)
    _style_header(ws, ["Website Name", "URL", "Active", "Schedule", "Threshold", "Desktop Score", "Mobile Score",
                       "Desktop LCP", "Mobile LCP", "Desktop CLS", "Mobile CLS", "Status", "Last Checked"],
                  [26, 48, 8, 34, 10, 12, 12, 12, 12, 11, 11, 11, 18])
    for w in snap.websites:
        _put(ws, [w["name"], w["url"], "Yes" if w["active"] else "No", w["schedule"], w["threshold"],
                   w["desktop_score"], w["mobile_score"], w["desktop_lcp"], w["mobile_lcp"],
                   w["desktop_cls"], w["mobile_cls"], w["status"], w["last_checked"]])
        fill = _FILLS.get(w["status"])
        if fill:
            ws.cell(ws.max_row, 12).fill = fill
        for col in (8, 9):
            ws.cell(ws.max_row, col).number_format = '0.00" s"'
    _put(ws, [])
    _put(ws, [f"Generated {gen} ({snap.tz_name})"])

    ws = wb.create_sheet(ATTENTION)
    _style_header(ws, ["Website Name", "URL", "Device", "Issue", "Performance Score", "Threshold", "FCP", "LCP",
                       "TBT", "CLS", "Speed Index", "Tested At", "Error", "PageSpeed Report"],
                  [26, 48, 9, 12, 12, 10, 9, 9, 10, 8, 12, 18, 40, 60])
    for a in snap.attention:
        _put(ws, [a["name"], a["url"], a["device"], a["issue"], a["score"], a["threshold"], a["fcp"], a["lcp"],
                   a["tbt"], a["cls"], a["si"], a["tested_at"], a["error"], a["link"]])
        ws.cell(ws.max_row, 4).fill = _FILLS["FAILED" if a["issue"] == "FAILED" else "ATTENTION"]
    if not snap.attention:
        _put(ws, ["No pages currently require attention."])
    _put(ws, [])
    _put(ws, [f"Generated {gen} ({snap.tz_name})"])

    ws = wb.create_sheet(SUMMARY)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    _put(ws, ["Monitoring Summary"])
    ws["A1"].font = Font(bold=True, size=14)
    for k, v in snap.totals.items():
        _put(ws, [k, v])
    _put(ws, [f"Generated {gen} ({snap.tz_name})"])
    _put(ws, [])
    start = ws.max_row + 1
    headers = ["Date", "Tests", "Successful", "Failed", "Avg Desktop Score", "Avg Mobile Score",
               "Below Threshold"]
    _put(ws, headers)
    for i in range(1, len(headers) + 1):
        c = ws.cell(start, i)
        c.fill, c.font = _HEADER_FILL, _HEADER_FONT
        ws.column_dimensions[get_column_letter(i)].width = max(ws.column_dimensions[get_column_letter(i)].width or 0, 16)
    for d in snap.daily:
        _put(ws, [d["date"], d["tests"], d["ok"], d["failed"], d["avg_desktop"], d["avg_mobile"], d["below"]])
        ws.cell(ws.max_row, 1).number_format = "DD-MMM-YYYY"


class WorkbookUnreadable(Exception):
    """The stored workbook cannot be opened (corrupt file / missing history sheet)."""


def _history_sheets(wb) -> list:
    names = [n for n in wb.sheetnames if n == HISTORY or (n.startswith(HISTORY + " ") and n[len(HISTORY) + 1:].isdigit())]
    return sorted((wb[n] for n in names), key=lambda w: 1 if w.title == HISTORY else int(w.title.rsplit(" ", 1)[1]))


def _existing_ids(sheets) -> set[int]:
    ids: set[int] = set()
    for ws in sheets:
        for (val,) in ws.iter_rows(min_row=2, min_col=RESULT_ID_COL, max_col=RESULT_ID_COL, values_only=True):
            if isinstance(val, (int, float)):
                ids.add(int(val))
    return ids


def _data_rows(sheets) -> int:
    return sum(max(0, ws.max_row - 1) for ws in sheets)


def _new_history_sheet(wb, index: int):
    title = HISTORY if index == 1 else f"{HISTORY} {index}"
    ws = wb.create_sheet(title, index - 1)
    _style_header(ws, HISTORY_HEADERS, _WIDTHS)
    return ws


def _open(path: Path):
    """Open the stored workbook; only failures *here* mean the file itself is unreadable."""
    try:
        wb = load_workbook(path)
    except Exception as exc:  # zipfile/XML errors of many types
        raise WorkbookUnreadable(f"{type(exc).__name__}: {exc}") from exc
    if not _history_sheets(wb):
        raise WorkbookUnreadable(f"Workbook has no '{HISTORY}' sheet")
    return wb


def _build(path: Path, out: Path, exists: bool, rows: list[tuple[int, list]], snap: Snapshot
           ) -> tuple[list[int], int, int]:
    """Runs in a worker thread. Returns (ids_written_or_present, total_data_rows, previous_data_rows).

    Raises WorkbookUnreadable only when the existing file cannot be opened; any other
    error (bad data, disk) propagates as-is and never triggers a rebuild.
    """
    if exists:
        wb = _open(path)
        sheets = _history_sheets(wb)
        present = _existing_ids(sheets)
    else:
        wb = Workbook()
        wb.remove(wb.active)
        sheets = [_new_history_sheet(wb, 1)]
        present = set()
    prev = _data_rows(sheets)
    ws = sheets[-1]
    done: list[int] = []
    for rid, values in rows:
        if rid not in present:
            if ws.max_row >= MAX_SHEET_ROWS:
                ws = _new_history_sheet(wb, len(sheets) + 1)
                sheets.append(ws)
            _append_history(ws, values)
            present.add(rid)
        done.append(rid)
    _write_snapshot_sheets(wb, snap)
    wb.active = 0
    wb.save(out)
    wb.close()
    # Verify the saved file before it may replace the stored copy.
    check = load_workbook(out, read_only=True)
    try:
        total = _data_rows(_history_sheets(check))
    finally:
        check.close()
    if total < prev:
        raise ExcelError(f"Refusing to save: history would shrink from {prev} to {total} rows")
    return done, total, prev


class ExcelReportService:
    def __init__(self, sf: async_sessionmaker[AsyncSession], storage: StorageService):
        self.sf = sf
        self.storage = storage
        self.filename = get_settings().excel_filename

    async def _snapshot(self, session: AsyncSession, tz: ZoneInfo, tz_name: str) -> Snapshot:
        now = now_utc()
        ov = await reporting.overview(session, tz, now)
        app = await settings_service.load(session)
        schedule_text = global_schedule.compute(app, now).description
        websites, attention = [], []

        def fmt(dt):
            return dt.astimezone(tz).strftime("%d-%b-%Y %H:%M") if dt else None

        for row in ov.rows:
            w: MonitoredWebsite = row["website"]
            d, m = row["desktop"], row["mobile"]
            websites.append({
                "name": w.name, "url": w.url, "active": w.is_active,
                "schedule": schedule_text if w.is_active else "Monitoring disabled",
                "threshold": w.threshold,
                "desktop_score": d.performance_score if d else None,
                "mobile_score": m.performance_score if m else None,
                "desktop_lcp": d.lcp_s if d else None, "mobile_lcp": m.lcp_s if m else None,
                "desktop_cls": d.cls if d else None, "mobile_cls": m.cls if m else None,
                "status": row["status"], "last_checked": fmt(row["last_checked"]),
            })
            if not w.is_active:
                continue
            for r in (d, m):
                if r is None:
                    continue
                st = reporting.result_status(r)
                if r.status == "success" and r.performance_score is not None and r.performance_score >= w.threshold:
                    continue
                attention.append({
                    "name": w.name, "url": w.url, "device": r.strategy.capitalize(),
                    "issue": "FAILED" if st == "FAILED" else "ATTENTION",
                    "score": r.performance_score, "threshold": w.threshold, "fcp": r.fcp_s, "lcp": r.lcp_s,
                    "tbt": r.tbt_ms, "cls": r.cls, "si": r.speed_index_s, "tested_at": fmt(r.tested_at),
                    "error": r.error_message if r.status != "success" else None,
                    "link": psi_report_link(w.url, r.strategy),
                })
        # Daily aggregates (all history, newest first).
        agg: dict = defaultdict(lambda: {"tests": 0, "ok": 0, "failed": 0, "d": [], "m": [], "below": 0})
        res = await session.execute(select(
            PerformanceResult.tested_at, PerformanceResult.strategy, PerformanceResult.status,
            PerformanceResult.performance_score, PerformanceResult.threshold))
        total_tests = ok_tests = 0
        for tested_at, strategy, status, score, threshold in res:
            day = tested_at.astimezone(tz).date()
            a = agg[day]
            a["tests"] += 1
            total_tests += 1
            if status == "success" and score is not None:
                a["ok"] += 1
                ok_tests += 1
                (a["d"] if strategy == "desktop" else a["m"]).append(score)
                if score < threshold:
                    a["below"] += 1
            else:
                a["failed"] += 1
        daily = [{
            "date": day, "tests": a["tests"], "ok": a["ok"], "failed": a["failed"],
            "avg_desktop": round(sum(a["d"]) / len(a["d"]), 1) if a["d"] else None,
            "avg_mobile": round(sum(a["m"]) / len(a["m"]), 1) if a["m"] else None,
            "below": a["below"],
        } for day, a in sorted(agg.items(), reverse=True)]
        totals = {
            "Total Websites": ov.total_websites,
            "Active Websites": ov.active_websites,
            "Total Tests Recorded": total_tests,
            "Successful Tests": ok_tests,
            "Failed Tests": total_tests - ok_tests,
            "Tests Today": ov.tests_today,
            "Average Desktop Performance (latest)": ov.avg_desktop,
            "Average Mobile Performance (latest)": ov.avg_mobile,
            "Pages Requiring Attention": len(attention),
        }
        return Snapshot(tz_name, now, websites, attention, daily, totals)

    async def _rows(self, session: AsyncSession, tz: ZoneInfo, pending_only: bool) -> list[tuple[int, list]]:
        q = select(PerformanceResult).order_by(PerformanceResult.tested_at, PerformanceResult.id)
        if pending_only:
            q = q.where(PerformanceResult.excel_appended.is_(False))
        return [(r.id, history_row(r, tz)) for r in (await session.scalars(q)).all()]

    async def sync(self) -> SyncResult:
        """Append every not-yet-exported result and refresh the snapshot sheets."""
        async with _process_lock:
            async with locks.lease(self.sf, LOCK_NAME, ttl_seconds=900, wait_seconds=600) as got:
                if not got:
                    raise ExcelError("Workbook is locked by another process; try again shortly")
                keeper = asyncio.create_task(self._keep_lease(got))
                try:
                    result = await self._sync_locked()
                except Exception as exc:
                    await self._record(False, f"{type(exc).__name__}: {exc}")
                    raise
                finally:
                    keeper.cancel()
                await self._record(True, f"Appended {result.appended} row(s); {result.total_rows - 1} total"
                                   + (" (rebuilt from database)" if result.rebuilt else ""))
                return result

    async def _keep_lease(self, owner: str) -> None:
        while True:
            await asyncio.sleep(120)
            try:
                await locks.refresh(self.sf, LOCK_NAME, 900, owner)
            except Exception:  # noqa: BLE001 - retried on the next tick
                log.warning("Could not refresh Excel lease", exc_info=True)

    async def _sync_locked(self) -> SyncResult:
        async with self.sf() as session:
            app = await settings_service.load(session)
            tz = app.tz
            snap = await self._snapshot(session, tz, app.timezone)
        with tempfile.TemporaryDirectory() as tmp:
            src, out = Path(tmp) / "current.xlsx", Path(tmp) / "next.xlsx"
            exists = await self.storage.download(self.filename, src)
            corrupt_copy = None
            async with self.sf() as session:
                rows = await self._rows(session, tz, pending_only=exists)
            try:
                done, total, prev = await asyncio.to_thread(_build, src, out, exists, rows, snap)
            except WorkbookUnreadable as exc:
                # Existing workbook unreadable: keep it, rebuild complete history from DB.
                stamp = now_utc().strftime("%Y%m%dT%H%M%SZ")
                corrupt_copy = self.filename.replace(".xlsx", f".corrupt-{stamp}.xlsx")
                log.error("Workbook could not be opened (%s); preserving as %s and rebuilding", exc, corrupt_copy)
                await self.storage.upload(src, corrupt_copy)
                async with self.sf() as session:
                    rows = await self._rows(session, tz, pending_only=False)
                done, total, prev = await asyncio.to_thread(_build, src, out, False, rows, snap)
                exists = False
            await self.storage.upload(out, self.filename)
        if done:
            async with self.sf() as session:
                now = now_utc()
                for i in range(0, len(done), 500):
                    await session.execute(
                        update(PerformanceResult)
                        .where(PerformanceResult.id.in_(done[i:i + 500]), PerformanceResult.excel_appended.is_(False))
                        .values(excel_appended=True, excel_appended_at=now)
                    )
                await session.commit()
        return SyncResult(appended=max(0, total - prev), rebuilt=not exists, total_rows=total + 1,
                          preserved_corrupt_copy=corrupt_copy)

    async def _record(self, ok: bool, message: str) -> None:
        try:
            async with self.sf() as session:
                await session.execute(
                    update(ScheduledTask).where(ScheduledTask.name == LOCK_NAME).values(
                        last_completed_at=now_utc(), last_status="ok" if ok else "error", last_message=message[:2000]))
                await session.commit()
        except Exception:  # pragma: no cover - logging only
            log.exception("Could not record Excel sync status")

    async def fetch_latest(self, dest: Path) -> bool:
        """Copy the latest stored workbook to dest (for downloads and e-mail attachments)."""
        return await self.storage.download(self.filename, dest)

    async def status(self) -> dict:
        async with self.sf() as session:
            row = await session.scalar(select(ScheduledTask).where(ScheduledTask.name == LOCK_NAME))
            pending = await session.scalar(
                select(PerformanceResult.id).where(PerformanceResult.excel_appended.is_(False)).limit(1))
        return {
            "last_updated_at": row.last_completed_at if row else None,
            "last_status": row.last_status if row else None,
            "last_message": row.last_message if row else None,
            "pending_rows": pending is not None,
        }
