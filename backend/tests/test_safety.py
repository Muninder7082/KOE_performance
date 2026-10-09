"""Excel safety, alert modes, security and real SMTP delivery."""
from __future__ import annotations

import asyncio
import email
import socket
import threading

from openpyxl import load_workbook
from sqlalchemy import update

from app import clock
from app.models import PerformanceResult
from tests.conftest import SCHED_HEADERS, ist, login, running_app, wait_job


def _rows(e):
    wb = load_workbook(e.storage_dir / "website-performance.xlsx")
    rows = list(wb["Performance History"].iter_rows(min_row=2, values_only=True))
    wb.close()
    return rows


async def _site_and_test(client, url="https://a.example.com/", strategies=("desktop", "mobile")):
    r = await client.post("/api/websites", json={"name": "A", "url": url})
    sid = r.json()["id"] if r.status_code == 201 else (await client.get("/api/websites")).json()["items"][0]["id"]
    r = await client.post(f"/api/websites/{sid}/test", json={"strategies": list(strategies)})
    job = await wait_job(client, r.json()["id"])
    assert job["status"] == "completed", job
    return sid


# ---- Excel -----------------------------------------------------------------------------
async def test_missing_workbook_is_rebuilt_with_full_history(env):
    async with running_app(env) as (app, client):
        await login(client)
        sid = await _site_and_test(client)
        await _site_and_test(client)
        assert len(_rows(env)) == 4
        (env.storage_dir / "website-performance.xlsx").unlink()
        (env.storage_dir / "website-performance.xlsx.bak").unlink()
        r = await client.post("/api/excel/sync")
        assert r.json()["rebuilt"] is True and r.json()["total_rows"] == 4
        assert sorted(row[18] for row in _rows(env)) == [1, 2, 3, 4]  # chronological order
        _ = sid


async def test_corrupt_workbook_is_preserved_never_overwritten(env):
    async with running_app(env) as (app, client):
        await login(client)
        await _site_and_test(client)
        path = env.storage_dir / "website-performance.xlsx"
        path.write_bytes(b"this is not a zip file")
        await _site_and_test(client)
        corrupt = list(env.storage_dir.glob("website-performance.corrupt-*.xlsx"))
        assert len(corrupt) == 1 and corrupt[0].read_bytes() == b"this is not a zip file"
        assert len(_rows(env)) == 4  # complete history rebuilt from the database


async def test_no_duplicate_rows_if_marking_failed(env):
    async with running_app(env) as (app, client):
        await login(client)
        await _site_and_test(client)
        # Simulate a crash after upload but before results were flagged as appended.
        async with _sf(app)() as s:
            await s.execute(update(PerformanceResult).values(excel_appended=False))
            await s.commit()
        await client.post("/api/excel/sync")
        assert len(_rows(env)) == 2


async def test_history_rows_never_modified(env):
    async with running_app(env) as (app, client):
        await login(client)
        await _site_and_test(client)
        before = _rows(env)
        env.psi.default = 50
        await _site_and_test(client)
        after = _rows(env)
        assert after[:2] == before and len(after) == 4


def _sf(app):
    from app import db

    return db.session_factory()


# ---- Alert modes ------------------------------------------------------------------------
async def test_alert_once_until_recovered_day_sequence(env):
    e = env
    seq = {1: 82, 2: 82, 3: 94, 4: 85}
    clock.set_now(ist(2026, 10, 1, 8, 0))  # before start-up: the schedule is anchored at install time
    async with running_app(e) as (app, client):
        await login(client)
        await client.post("/api/websites", json={"name": "A", "url": "https://a.example.com/"})
        for day, score in seq.items():
            e.psi.scores[("https://a.example.com/", "mobile")] = score
            clock.set_now(ist(2026, 10, 1 + day, 8, 45))
            r = await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
            assert r.json()["websites_due"] == 1
        alerts = [m for m in e.outbox.sent if m["type"] == "performance_alert"]
        assert len(alerts) == 2  # day 1 and day 4 only
        assert len([m for m in e.outbox.sent if m["type"] == "recovery_notification"]) == 1
        await login(client)
        items = (await client.get("/api/alerts")).json()["items"]
        assert [a["state"] for a in items] == ["open", "recovered"]


async def test_alert_mode_daily_and_every(env):
    e = env
    e.psi.default = 60
    async with running_app(e) as (app, client):
        clock.set_now(ist(2026, 10, 1, 8, 0))
        await login(client)
        r = await client.put("/api/settings", json={"alert_mode": "daily"})
        assert r.status_code == 200
        sid = (await client.post("/api/websites", json={"name": "A", "url": "https://a.example.com/"})).json()["id"]
        for _ in range(2):
            r = await client.post(f"/api/websites/{sid}/test", json={"strategies": ["mobile"]})
            await wait_job(client, r.json()["id"])
        assert len([m for m in e.outbox.sent if m["type"] == "performance_alert"]) == 1
        clock.set_now(ist(2026, 10, 2, 9, 0))
        await login(client)
        r = await client.post(f"/api/websites/{sid}/test", json={"strategies": ["mobile"]})
        await wait_job(client, r.json()["id"])
        assert len([m for m in e.outbox.sent if m["type"] == "performance_alert"]) == 2
        await client.put("/api/settings", json={"alert_mode": "every_occurrence"})
        r = await client.post(f"/api/websites/{sid}/test", json={"strategies": ["mobile"]})
        await wait_job(client, r.json()["id"])
        assert len([m for m in e.outbox.sent if m["type"] == "performance_alert"]) == 3


async def test_per_site_threshold(env):
    e = env
    e.psi.default = 85
    async with running_app(e) as (app, client):
        await login(client)
        sid = (await client.post("/api/websites", json={"name": "A", "url": "https://a.example.com/",
                                                         "threshold": 80})).json()["id"]
        r = await client.post(f"/api/websites/{sid}/test", json={"strategies": ["mobile"]})
        await wait_job(client, r.json()["id"])
        assert not [m for m in e.outbox.sent if m["type"] == "performance_alert"]
        assert (await client.get(f"/api/websites/{sid}")).json()["status"] == "GOOD"


# ---- Security ---------------------------------------------------------------------------
async def test_auth_csrf_roles_and_rate_limits(env):
    async with running_app(env) as (app, client):
        assert (await client.get("/api/websites")).status_code == 401
        assert (await client.get("/api/excel/download")).status_code == 401
        r = await client.post("/api/auth/login", json={"email": "admin@example.com", "password": "wrong-password"})
        assert r.status_code == 401
        await login(client)
        token = client.headers.pop("X-CSRF-Token")
        r = await client.post("/api/websites", json={"name": "A", "url": "https://a.example.com/"})
        assert r.status_code == 403 and "CSRF" in r.json()["detail"]
        client.headers["X-CSRF-Token"] = token
        r = await client.post("/api/users", json={"email": "viewer@example.com", "password": "Viewer-Pass-123",
                                                  "role": "viewer"})
        assert r.status_code == 201
        # Viewer: read-only
        await client.post("/api/auth/logout")
        await login(client, "viewer@example.com", "Viewer-Pass-123")
        assert (await client.get("/api/websites")).status_code == 200
        assert (await client.post("/api/websites", json={"name": "A", "url": "https://a.example.com/"})).status_code == 403
        assert (await client.put("/api/settings", json={"default_threshold": 50})).status_code == 403
        # Settings never expose secrets
        body = (await client.get("/api/settings")).text
        for secret in ("TEST-PSI-KEY-123", "test-scheduler-token", "test-secret-key"):
            assert secret not in body
        # Login brute force throttled
        await client.post("/api/auth/logout")
        codes = [(await client.post("/api/auth/login", json={"email": "x@example.com", "password": "nope"})).status_code
                 for _ in range(12)]
        assert codes.count(429) >= 1
        # Security headers
        r = await client.get("/api/health")
        assert "frame-ancestors" in r.headers["content-security-policy"]
        assert r.headers["x-content-type-options"] == "nosniff"


async def test_password_change_revokes_sessions(env):
    async with running_app(env) as (app, client):
        await login(client)
        old_cookie = client.cookies.get("pt_session")
        r = await client.post("/api/auth/change-password",
                              json={"current_password": "Admin-Password-123", "new_password": "New-Password-456"})
        assert r.status_code == 200
        client.headers["X-CSRF-Token"] = r.json()["csrf_token"]
        assert (await client.get("/api/auth/me")).status_code == 200
        client.cookies.set("pt_session", old_cookie)
        assert (await client.get("/api/auth/me")).status_code == 401


async def test_input_validation_and_xss_escaping(env):
    e = env
    e.psi.default = 50
    async with running_app(e) as (app, client):
        await login(client)
        bad = [{"name": "A", "url": "https://a.example.com/", "monitor_time": "25:00"},
               {"name": "A", "url": "https://a.example.com/", "frequency": "yearly"},
               {"name": "A", "url": "https://a.example.com/", "threshold": 0},
               {"name": "", "url": "https://a.example.com/"},
               {"name": "A", "url": "https://a.example.com/", "timezone": "Mars/Base"}]
        for body in bad:
            assert (await client.post("/api/websites", json=body)).status_code == 422, body
        evil = '<script>alert("x")</script>'
        sid = (await client.post("/api/websites", json={"name": evil, "url": "https://a.example.com/"})).json()["id"]
        r = await client.post(f"/api/websites/{sid}/test", json={"strategies": ["mobile"]})
        await wait_job(client, r.json()["id"])
        alert = [m for m in e.outbox.sent if m["type"] == "performance_alert"][0]
        assert evil not in alert["html"] and "&lt;script&gt;" in alert["html"]
        # LIKE wildcards are escaped in search
        assert (await client.get("/api/websites", params={"search": "%"})).json()["total"] == 0


# ---- Real SMTP delivery through a local SMTP server ------------------------------------------
class MiniSMTP(threading.Thread):
    def __init__(self):
        super().__init__(daemon=True)
        self.sock = socket.socket()
        self.sock.bind(("127.0.0.1", 0))
        self.sock.listen(1)
        self.port = self.sock.getsockname()[1]
        self.messages: list[bytes] = []

    def run(self):
        conn, _ = self.sock.accept()
        f = conn.makefile("rb")
        conn.sendall(b"220 mini ESMTP\r\n")
        while True:
            line = f.readline()
            if not line:
                break
            cmd = line.strip().upper()
            if cmd.startswith(b"EHLO") or cmd.startswith(b"HELO"):
                conn.sendall(b"250-mini\r\n250 SIZE 52428800\r\n")
            elif cmd == b"DATA":
                conn.sendall(b"354 go\r\n")
                data = []
                while True:
                    ln = f.readline()
                    if ln in (b".\r\n", b".\n"):
                        break
                    data.append(ln)
                self.messages.append(b"".join(data))
                conn.sendall(b"250 queued\r\n")
            elif cmd == b"QUIT":
                conn.sendall(b"221 bye\r\n")
                break
            else:
                conn.sendall(b"250 ok\r\n")
        conn.close()


async def test_smtp_daily_report_with_attachment(env, monkeypatch):
    from app.config import get_settings
    from app.services.email_service import EmailService

    server = MiniSMTP()
    server.start()
    monkeypatch.setenv("SMTP_HOST", "127.0.0.1")
    monkeypatch.setenv("SMTP_PORT", str(server.port))
    monkeypatch.setenv("SMTP_SECURITY", "none")
    get_settings.cache_clear()
    async with running_app(env) as (app, client):
        await login(client)
        await _site_and_test(client)
        monkeypatch.undo()  # restore real _deliver for the report only
        monkeypatch.setenv("SMTP_HOST", "127.0.0.1")
        monkeypatch.setenv("SMTP_PORT", str(server.port))
        monkeypatch.setenv("SMTP_SECURITY", "none")
        get_settings.cache_clear()
        from app.container import get_container

        c = get_container()
        c.email.s = get_settings()
        assert EmailService._deliver.__name__ == "_deliver"
        ok = await c.report.send_daily()
        assert ok
        await asyncio.to_thread(server.join, 5)
        msg = email.message_from_bytes(server.messages[0])
        assert msg["Subject"].startswith("Daily Website Performance Report - ")
        assert msg["To"] == "reports@example.com"
        parts = {p.get_filename(): p for p in msg.walk() if p.get_filename()}
        payload = parts["website-performance.xlsx"].get_payload(decode=True)
        import io

        assert load_workbook(io.BytesIO(payload))["Performance History"].max_row == 3


async def test_deleting_website_closes_open_alerts_and_keeps_history(env):
    env.psi.default = 60
    async with running_app(env) as (app, client):
        await login(client)
        sid = await _site_and_test(client, strategies=("mobile",))
        assert (await client.get("/api/alerts", params={"state": "open"})).json()["total"] == 1
        assert (await client.delete(f"/api/websites/{sid}")).status_code == 200
        assert (await client.get("/api/alerts", params={"state": "open"})).json()["total"] == 0
        assert (await client.get("/api/alerts", params={"state": "closed"})).json()["total"] == 1
        assert (await client.get("/api/results")).json()["items"][0]["website_name"] == "A"


async def test_forged_forwarded_for_cannot_bypass_rate_limit(env):
    async with running_app(env) as (app, client):
        codes = []
        for i in range(12):  # attacker rotates the left-most entry; the proxy appends the real IP
            r = await client.post("/api/auth/login", json={"email": f"u{i}@example.com", "password": "x"},
                                  headers={"X-Forwarded-For": f"10.0.0.{i}, 198.51.100.7"})
            codes.append(r.status_code)
        assert 429 in codes


# ---- Regression tests for the independent review findings ------------------------------
async def test_control_characters_do_not_break_excel(env):
    env.psi.scores[("https://a.example.com/", "mobile")] = ("http", 500, {"error": {"message": "Lighthouse returned error: bad\x1b[31m\x0bchars"}})
    async with running_app(env) as (app, client):
        await login(client)
        sid = (await client.post("/api/websites", json={"name": "Bad\x0bName", "url": "https://a.example.com/"})).json()["id"]
        r = await client.post(f"/api/websites/{sid}/test", json={"strategies": ["desktop", "mobile"]})
        job = await wait_job(client, r.json()["id"])
        assert job["warning"] is None, job
        rows = _rows(env)
        assert len(rows) == 2 and rows[0][2] == "BadName"
        assert not list(env.storage_dir.glob("*.corrupt-*"))


async def test_history_rolls_over_to_new_sheet(env, monkeypatch):
    from app.services import excel as excel_mod

    monkeypatch.setattr(excel_mod, "MAX_SHEET_ROWS", 3)  # header + 2 rows per sheet
    async with running_app(env) as (app, client):
        await login(client)
        await _site_and_test(client)
        await _site_and_test(client)
        await _site_and_test(client)
        wb = load_workbook(env.storage_dir / "website-performance.xlsx")
        names = wb.sheetnames
        counts = [wb[n].max_row - 1 for n in names if n.startswith("Performance History")]
        ids = sorted(r[18] for n in names if n.startswith("Performance History")
                     for r in wb[n].iter_rows(min_row=2, values_only=True))
        assert names[:3] == ["Performance History", "Performance History 2", "Performance History 3"]
        assert counts == [2, 2, 2] and ids == [1, 2, 3, 4, 5, 6]


async def test_restart_marks_interrupted_runs_aborted(env):
    from app import db
    from app.models import TaskRun

    async with running_app(env) as (app, client):
        async with db.session_factory()() as s:
            s.add(TaskRun(run_id="run-crashed", task="scheduled_checks", trigger="external", status="running",
                          started_at=clock.now_utc()))
            await s.commit()
    async with running_app(env) as (app, client):
        await login(client)
        runs = (await client.get("/api/task-runs")).json()["items"]
        assert runs[0]["run_id"] == "run-crashed" and runs[0]["status"] == "aborted"


async def test_failed_daily_report_is_retried(env):
    e = env
    clock.set_now(ist(2026, 10, 8, 8, 0))
    async with running_app(e) as (app, client):
        e.outbox.fail = True
        clock.set_now(ist(2026, 10, 8, 8, 41))
        r = await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
        assert r.json()["report_sent"] is False
        e.outbox.fail = False
        clock.set_now(ist(2026, 10, 8, 8, 51))
        r = await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
        assert r.json()["report_sent"] is True
        clock.set_now(ist(2026, 10, 8, 9, 1))
        r = await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
        assert r.json()["report_sent"] is None  # sent once, not again
        assert len([m for m in e.outbox.sent if m["type"] == "daily_report"]) == 1


async def test_occurrence_claim_is_exclusive(env):
    clock.set_now(ist(2026, 10, 8, 8, 0))
    async with running_app(env) as (app, client):
        await login(client)
        sid = (await client.post("/api/websites", json={"name": "A", "url": "https://a.example.com/"})).json()["id"]
        from app.container import get_container

        sched = get_container().scheduler
        occ = ist(2026, 10, 8, 8, 40)
        results = await asyncio.gather(*(sched._claim(sid, occ) for _ in range(5)))
        assert results.count(True) == 1


async def test_saved_lighthouse_report_is_served_and_purged(env, monkeypatch):
    e = env
    e.psi.scores[("https://b.example.com/", "mobile")] = ("http", 500, {"error": {"message": "Lighthouse returned error: FAILED_DOCUMENT_REQUEST"}})
    clock.set_now(ist(2026, 10, 1, 9, 0))
    async with running_app(e) as (app, client):
        await login(client)
        await _site_and_test(client, strategies=("desktop", "mobile"))
        items = (await client.get("/api/results")).json()["items"]
        assert all(i["has_report"] for i in items)
        rid = items[0]["id"]
        r = await client.get(f"/api/results/{rid}/report")
        assert r.status_code == 200 and "text/html" in r.headers["content-type"]
        assert 'id="lhr-json"' in r.text and "/lighthouse-assets/standalone.js" in r.text
        other = [i for i in items if i["id"] != rid][0]
        assert 'class="tab active"' in r.text and f'href="/api/results/{other["id"]}/report"' in r.text  # Mobile/Desktop tabs
        assert "<script>" not in r.text  # no inline JS: CSP stays script-src 'self'
        assert "script-src 'self'" in r.headers["content-security-policy"]
        assert (await client.get("/lighthouse-assets/standalone.js")).status_code == 200
        j = await client.get(f"/api/results/{rid}/viewer-url")
        assert j.status_code == 200 and j.json()["url"].startswith("https://googlechrome.github.io/lighthouse/viewer/?gzip=1#")
        v = await client.get(f"/api/results/{rid}/viewer")
        assert v.status_code == 200 and "googlechrome.github.io/lighthouse/viewer/?gzip=1#" in v.text
        assert "<script>" not in v.text and (await client.get("/lighthouse-assets/open-viewer.js")).status_code == 200
        import base64, gzip as _gz, json as _json, re
        frag = re.search(r"#([A-Za-z0-9+/=]+)", v.text).group(1)
        assert _json.loads(_gz.decompress(base64.b64decode(frag)))["lhr"]["requestedUrl"] == "https://a.example.com/"
        assert len(list((e.storage_dir / "reports").rglob("*.json.gz"))) == 2
        # failed test: no report -> friendly 404 page
        sid2 = (await client.post("/api/websites", json={"name": "B", "url": "https://b.example.com/"})).json()["id"]
        r = await client.post(f"/api/websites/{sid2}/test", json={"strategies": ["mobile"]})
        await wait_job(client, r.json()["id"])
        failed = (await client.get("/api/results", params={"status": "failed"})).json()["items"][0]
        assert failed["has_report"] is False
        r = await client.get(f"/api/results/{failed['id']}/report")
        assert r.status_code == 404 and "Report not available" in r.text
        assert (await client.get(f"/api/results/{failed['id']}/viewer-url")).status_code == 404
        # unauthenticated access is refused
        client.cookies.clear()
        assert (await client.get(f"/api/results/{rid}/report")).status_code == 401
    # 100 days later the scheduled cycle removes reports older than 90 days (scores stay)
    clock.set_now(ist(2027, 1, 10, 9, 0))
    async with running_app(e) as (app, client):
        await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
        assert not list((e.storage_dir / "reports" / "2026").rglob("*.json.gz"))  # old ones removed
        await login(client)
        old = (await client.get("/api/results", params={"date_to": "2026-12-31"})).json()["items"]
        assert len(old) == 3 and not any(i["has_report"] for i in old)  # scores kept, reports gone
        new = (await client.get("/api/results", params={"date_from": "2027-01-01"})).json()["items"]
        assert new and all(i["has_report"] for i in new if i["status"] == "success")


async def test_brevo_provider_sends_over_https(env, monkeypatch):
    import base64 as _b64
    import json as _json

    import httpx

    from app.config import get_settings
    from app.services.email_service import Attachment, EmailService, OutgoingEmail

    monkeypatch.setenv("EMAIL_PROVIDER", "brevo")
    monkeypatch.setenv("BREVO_API_KEY", "xkeysib-test-123")
    get_settings.cache_clear()
    seen = {}

    def handler(request: httpx.Request) -> httpx.Response:
        seen["url"] = str(request.url)
        seen["key"] = request.headers.get("api-key")
        seen["body"] = _json.loads(request.content)
        return httpx.Response(201, json={"messageId": "<x@brevo>"})

    async with running_app(env) as (app, client):
        from app.container import get_container

        sf = get_container().sf
        monkeypatch.undo()  # restore the real _deliver patched by the env fixture
        monkeypatch.setenv("EMAIL_PROVIDER", "brevo")
        monkeypatch.setenv("BREVO_API_KEY", "xkeysib-test-123")
        get_settings.cache_clear()
        svc = EmailService(sf, httpx.AsyncClient(transport=httpx.MockTransport(handler)))
        ok = await svc.send(OutgoingEmail("test_email", ["a@example.com", "b@example.com"], "Hello", "<p>Hi</p>", "Hi",
                                          attachments=[Attachment("website-performance.xlsx", b"PK-data")]))
        assert ok
        assert seen["url"] == "https://api.brevo.com/v3/smtp/email" and seen["key"] == "xkeysib-test-123"
        b = seen["body"]
        assert b["sender"]["email"] == "monitor@example.com" and [t["email"] for t in b["to"]] == ["a@example.com", "b@example.com"]
        assert b["attachment"][0]["name"] == "website-performance.xlsx"
        assert _b64.b64decode(b["attachment"][0]["content"]) == b"PK-data"
        assert svc.status()["configured"] is True


async def test_idle_scheduler_ticks_do_not_query_the_database(env):
    from sqlalchemy import event

    from app import db

    clock.set_now(ist(2026, 10, 8, 8, 0))
    async with running_app(env) as (app, client):
        await login(client)
        await client.post("/api/websites", json={"name": "A", "url": "https://a.example.com/"})
        clock.set_now(ist(2026, 10, 8, 8, 41))
        r = await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
        assert r.json()["websites_due"] == 1 and r.json()["report_sent"] is True
        queries = []
        listener = lambda *a, **k: queries.append(1)  # noqa: E731
        event.listen(db.get_engine().sync_engine, "before_cursor_execute", listener)
        try:
            for minute in (42, 43, 50, 59):
                clock.set_now(ist(2026, 10, 8, 8, minute))
                r = await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
                assert r.json()["status"] == "idle"
            clock.set_now(ist(2026, 10, 8, 23, 0))  # outside the catch-up window, nothing due
            assert (await client.post("/api/internal/run-scheduled-checks", headers=SCHED_HEADERS)).json()["status"] == "idle"
        finally:
            event.remove(db.get_engine().sync_engine, "before_cursor_execute", listener)
        assert queries == []
        # Changing the schedule is picked up on the next tick
        await login(client)
        assert (await client.put("/api/settings", json={"default_monitor_time": "23:30"})).status_code == 200
        clock.set_now(ist(2026, 10, 8, 23, 31))
        r = await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
        assert r.json()["websites_due"] == 1
