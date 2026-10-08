"""End-to-end workflow (spec section 35), driven through the HTTP API.

Add website -> configure daily 08:40 -> manual test (desktop + mobile) -> DB rows
-> Excel rows -> alert e-mail for score < 90 -> daily report with the complete
workbook attached -> scheduled run via the secure endpoint -> restart -> data and
Excel still there -> next day's run appends rows without removing any.
"""
from __future__ import annotations

import io

from openpyxl import load_workbook

from app import clock
from tests.conftest import SCHED_HEADERS, ist, login, running_app, wait_job

URL = "https://www.example.com/"
URL2 = "https://www.example.com/courses"


def _history(e):
    wb = load_workbook(e.storage_dir / "website-performance.xlsx")
    ws = wb["Performance History"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True)]
    sheets = wb.sheetnames
    wb.close()
    return rows, sheets


async def test_full_workflow(env):
    e = env
    clock.set_now(ist(2026, 10, 8, 8, 0))
    e.psi.scores[(URL, "desktop")] = 94
    e.psi.scores[(URL, "mobile")] = 82  # below threshold -> attention e-mail
    e.psi.scores[(URL2, "desktop")] = 96
    e.psi.scores[(URL2, "mobile")] = 93

    async with running_app(e) as (app, client):
        await login(client)
        # Add website with default schedule (daily 08:40 Asia/Kolkata, threshold 90)
        r = await client.post("/api/websites", json={"name": "Homepage", "url": "https://WWW.example.com"})
        assert r.status_code == 201, r.text
        site = r.json()
        assert site["frequency"] == "daily" and site["monitor_time"] == "08:40"
        assert site["timezone"] == "Asia/Kolkata" and site["threshold"] == 90
        sid = site["id"]
        # Duplicate URL rejected, invalid URL rejected
        assert (await client.post("/api/websites", json={"name": "Dup", "url": "http://www.example.com/"})).status_code == 409
        assert (await client.post("/api/websites", json={"name": "Bad", "url": "not a url"})).status_code == 422
        # Second page — every page follows the one schedule from Settings (daily 08:40)
        r = await client.post("/api/websites", json={"name": "Courses", "url": URL2})
        sid2 = r.json()["id"]

        # Manual test: both devices
        r = await client.post(f"/api/websites/{sid}/test", json={"strategies": ["desktop", "mobile"]})
        assert r.status_code == 202, r.text
        job = await wait_job(client, r.json()["id"])
        assert job["status"] == "completed", job
        assert job["stages"][:1] == ["Testing..."] and "Updating Excel..." in job["stages"]
        assert job["stages"][-1] == "Completed." and job["succeeded"] == 2

        # PostgreSQL/DB results
        res = (await client.get("/api/results", params={"website_id": sid})).json()
        assert res["total"] == 2
        by_dev = {x["strategy"]: x for x in res["items"]}
        assert by_dev["desktop"]["performance_score"] == 94 and by_dev["mobile"]["performance_score"] == 82
        assert by_dev["mobile"]["health"] == "ATTENTION" and all(x["excel_appended"] for x in res["items"])

        # Excel: two rows (Desktop + Mobile) and all four sheets
        rows, sheets = _history(e)
        assert sheets == ["Performance History", "Latest Status", "Attention Required", "Monitoring Summary"]
        assert [(r[2], r[4], r[5], r[14], r[16]) for r in rows] == [
            ("Homepage", "Desktop", 94, "GOOD", "SUCCESS"), ("Homepage", "Mobile", 82, "ATTENTION", "SUCCESS")]

        # Attention e-mail sent once (mobile 82 < 90)
        alerts = [m for m in e.outbox.sent if m["type"] == "performance_alert"]
        assert len(alerts) == 1
        assert alerts[0]["subject"] == "Performance Attention Required - Homepage"
        assert "Mobile" in alerts[0]["text"] and "82" in alerts[0]["text"] and "pagespeed.web.dev" in alerts[0]["text"]
        assert alerts[0]["to"] == ["alerts@example.com"]

        # Scheduled endpoint security
        assert (await client.post("/api/internal/run-scheduled-checks")).status_code == 401
        assert (await client.post("/api/internal/run-scheduled-checks",
                                  headers={"Authorization": "Bearer wrong"})).status_code == 401
        # Private-Space style: HF token in Authorization, scheduler token in X-Scheduler-Token
        r = await client.get("/api/internal/scheduler-status", headers={
            "Authorization": "Bearer hf_xxx", "X-Scheduler-Token": SCHED_HEADERS["Authorization"][7:]})
        assert r.status_code == 200 and r.json()["running"] is False

        # 08:41 IST: the common 08:40 run tests ALL pages, then e-mails the report
        clock.set_now(ist(2026, 10, 8, 8, 41))
        r = await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
        assert r.status_code == 200, r.text
        summary = r.json()
        assert summary["status"] == "completed" and summary["websites_due"] == 2
        assert summary["tests_succeeded"] == 4 and summary["report_sent"] is True
        # Second trigger in the same minute: nothing due, no duplicate report
        r = await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
        assert r.json()["websites_due"] == 0 and r.json()["report_sent"] is None

        rows, _ = _history(e)
        assert len(rows) == 6
        # Once-until-recovered: the repeated 82 did not send another alert
        assert len([m for m in e.outbox.sent if m["type"] == "performance_alert"]) == 1

        # Daily report with the complete historical workbook attached
        reports = [m for m in e.outbox.sent if m["type"] == "daily_report"]
        assert len(reports) == 1
        rep = reports[0]
        assert rep["subject"] == "Daily Website Performance Report - 08-Oct-2026"
        assert "Pages Requiring Attention" in rep["text"] and "Homepage" in rep["text"]
        (fname, content), = rep["attachments"]
        assert fname == "website-performance.xlsx"
        att = load_workbook(io.BytesIO(content))["Performance History"]
        assert att.max_row - 1 == 6  # all history rows, including this run's, are attached

        # A page added after today's start time waits for the next run
        await client.post("/api/websites", json={"name": "Late page", "url": "https://late.example.com/"})
        clock.set_now(ist(2026, 10, 8, 10, 1))
        r = await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
        assert r.json()["websites_due"] == 0
        late = (await client.get("/api/websites", params={"search": "Late"})).json()["items"][0]
        assert (await client.delete(f"/api/websites/{late['id']}")).status_code == 200

        # Logs available
        assert (await client.get("/api/email-logs")).json()["total"] >= 2
        runs = (await client.get("/api/task-runs")).json()
        assert runs["total"] == 3 and runs["items"][0]["status"] == "completed"
        dash = (await client.get("/api/dashboard")).json()
        assert dash["cards"]["total_websites"] == 2 and dash["cards"]["attention"] == 1
        assert dash["schedule"]["description"].startswith("Daily at 08:40 AM")
        assert dash["cards"]["tests_today"] == 6

    # ---- Restart the application: same database + same storage ----------------------
    clock.set_now(ist(2026, 10, 9, 8, 42))
    e.psi.scores[(URL, "mobile")] = 94  # recovered
    async with running_app(e) as (app, client):
        await login(client)
        sites = (await client.get("/api/websites")).json()
        assert sites["total"] == 2
        assert (await client.get("/api/results")).json()["total"] == 6
        assert len(_history(e)[0]) == 6  # Excel survived the restart

        # Next day's scheduled run appends WITHOUT deleting earlier rows
        r = await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
        assert r.json()["websites_due"] == 2 and r.json()["report_sent"] is True
        rows, _ = _history(e)
        assert len(rows) == 10
        dates = sorted({row[0].strftime("%d-%b-%Y") for row in rows})
        assert dates == ["08-Oct-2026", "09-Oct-2026"]
        assert [row[18] for row in rows] == sorted(row[18] for row in rows)  # appended in order, none removed

        # Recovery: alert closed + recovery e-mail
        alerts = (await client.get("/api/alerts")).json()["items"]
        assert alerts[0]["state"] == "recovered" and alerts[0]["recovered_score"] == 94
        assert any(m["type"] == "recovery_notification" for m in e.outbox.sent)

        # History endpoint (charts)
        hist = (await client.get(f"/api/websites/{sid}/history", params={"days": 7})).json()["items"]
        assert len(hist) == 6  # 2 manual + 2 + 2 scheduled

        # Excel download returns the full workbook
        r = await client.get("/api/excel/download")
        assert r.status_code == 200
        assert load_workbook(io.BytesIO(r.content))["Performance History"].max_row - 1 == 10

        # Deleting a website keeps its history (DB and Excel)
        assert (await client.delete(f"/api/websites/{sid2}")).status_code == 200
        assert (await client.get("/api/results")).json()["total"] == 10
        await client.post("/api/excel/sync")
        assert len(_history(e)[0]) == 10


async def test_one_failure_does_not_stop_others(env):
    e = env
    clock.set_now(ist(2026, 10, 8, 8, 0))
    urls = ["https://a.example.com/", "https://b.example.com/", "https://c.example.com/"]
    e.psi.scores[(urls[1], "desktop")] = ("http", 500, {"error": {"message": "Lighthouse returned error: FAILED_DOCUMENT_REQUEST"}})
    e.psi.scores[(urls[1], "mobile")] = ("timeout",)
    async with running_app(e) as (app, client):
        await login(client)
        for i, u in enumerate(urls):
            assert (await client.post("/api/websites", json={"name": f"Site {i}", "url": u})).status_code == 201
        clock.set_now(ist(2026, 10, 8, 8, 45))
        r = await client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS)
        s = r.json()
        assert s["websites_due"] == 3 and s["tests_succeeded"] == 4 and s["tests_failed"] == 2
        failed = (await client.get("/api/results", params={"status": "failed"})).json()["items"]
        assert {f["error_code"] for f in failed} == {"website_unavailable", "timeout"}
        rows, _ = _history(e)
        assert len(rows) == 6
        failed_rows = [r for r in rows if r[16] == "FAILED"]
        assert len(failed_rows) == 2 and all(r[14] == "FAILED" and r[17] for r in failed_rows)
        dash = (await client.get("/api/dashboard")).json()
        assert dash["cards"]["failing_websites"] == 1


async def test_background_trigger_and_concurrent_triggers(env):
    e = env
    clock.set_now(ist(2026, 10, 8, 8, 0))
    async with running_app(e) as (app, client):
        await login(client)
        await client.post("/api/websites", json={"name": "A", "url": "https://a.example.com/"})
        clock.set_now(ist(2026, 10, 8, 8, 41))
        import asyncio

        r1, r2 = await asyncio.gather(
            client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS),
            client.post("/api/internal/run-scheduled-checks?wait=true", headers=SCHED_HEADERS))
        statuses = sorted([r1.json()["status"], r2.json()["status"]])
        due_total = sum(r.json().get("websites_due", 0) for r in (r1, r2))
        assert due_total == 1 and "completed" in statuses  # the site is tested exactly once
        assert len(e.psi.calls) == 2
        clock.set_now(ist(2026, 10, 9, 8, 41))
        await login(client)  # the 12 h session from yesterday has expired
        r = await client.post("/api/internal/run-scheduled-checks", headers=SCHED_HEADERS)
        assert r.json()["status"] == "started" and r.json()["websites_due"] == 1
        for _ in range(200):
            d = (await client.get("/api/dashboard")).json()
            assert "scheduler_running" in d, d
            if len(e.psi.calls) == 4 and not d["scheduler_running"]:
                break
            await asyncio.sleep(0.05)
        assert len(e.psi.calls) == 4


async def test_manual_test_protection(env, monkeypatch):
    e = env
    monkeypatch.setenv("MANUAL_TEST_COOLDOWN_SECONDS", "120")
    from app.config import get_settings

    get_settings.cache_clear()
    async with running_app(e) as (app, client):
        await login(client)
        sid = (await client.post("/api/websites", json={"name": "A", "url": "https://a.example.com/"})).json()["id"]
        r = await client.post(f"/api/websites/{sid}/test", json={"strategies": ["mobile"]})
        await wait_job(client, r.json()["id"])
        r = await client.post(f"/api/websites/{sid}/test", json={"strategies": ["mobile"]})
        assert r.status_code == 429 and "wait" in r.json()["detail"]
        assert len(e.psi.calls) == 1
        r = await client.post(f"/api/websites/{sid}/test", json={"strategies": ["desktop"]})
        assert r.status_code == 202
        assert (await wait_job(client, r.json()["id"]))["succeeded"] == 1


async def test_email_failure_is_logged_and_alert_retried(env):
    e = env
    e.outbox.fail = True
    e.psi.default = 70
    async with running_app(e) as (app, client):
        await login(client)
        sid = (await client.post("/api/websites", json={"name": "A", "url": "https://a.example.com/"})).json()["id"]
        r = await client.post(f"/api/websites/{sid}/test", json={"strategies": ["mobile"]})
        assert (await wait_job(client, r.json()["id"]))["status"] == "completed"
        logs = (await client.get("/api/email-logs")).json()["items"]
        assert logs[0]["status"] == "failed" and "unreachable" in logs[0]["error"]
        e.outbox.fail = False
        r = await client.post(f"/api/websites/{sid}/test", json={"strategies": ["mobile"]})
        await wait_job(client, r.json()["id"])
        assert len([m for m in e.outbox.sent if m["type"] == "performance_alert"]) == 1
